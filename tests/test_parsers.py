import asyncio
import io

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image

from hybrid_rag.ingestion import pipeline
from hybrid_rag.ingestion.parser.office.parser import parse_docx
from hybrid_rag.ingestion.parser.spreadsheet.parser import parse_xlsx


def _png(path, size):
    Image.new("RGB", size, "white").save(path, "PNG")
    return path


def _docx(tmp_path):
    big, small = _png(tmp_path / "big.png", (900, 900)), _png(tmp_path / "small.png", (100, 100))
    doc = Document()
    doc.add_paragraph("Pengantar")
    doc.add_picture(str(small))
    doc.add_picture(str(big))
    path = tmp_path / "a.docx"
    doc.save(path)
    return path


def _xlsx(tmp_path):
    big, small = _png(tmp_path / "big.png", (900, 900)), _png(tmp_path / "small.png", (100, 100))
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    ws["A1"], ws["B1"] = "Nama", "Nilai"
    ws["A2"], ws["B2"] = "x", 1
    ws.add_image(XlImage(str(small)), "D2")
    ws.add_image(XlImage(str(big)), "D20")
    path = tmp_path / "a.xlsx"
    wb.save(path)
    return path


def _images(items):
    return [i for i in items if i["type"] == "image"]


def test_docx_keeps_only_large_images(tmp_path):
    items = parse_docx(_docx(tmp_path), tmp_path / "out")
    [image] = _images(items)
    assert Image.open(image["img_path"]).size == (900, 900)
    assert any(i["type"] == "text" and i["text"] == "Pengantar" for i in items)


def test_xlsx_keeps_only_large_images_and_uses_sheet_as_page(tmp_path):
    items = parse_xlsx(_xlsx(tmp_path), tmp_path / "out")
    [image] = _images(items)
    assert Image.open(image["img_path"]).size == (900, 900)
    assert image["page_idx"] == "Rekap"


def test_parse_file_hosts_figures_once_for_office_files(tmp_path, monkeypatch):
    hosted = []
    monkeypatch.setattr(pipeline, "host_figures", hosted.append)

    async def run(path):
        async with pipeline.parse_file(path) as content_list:
            return content_list

    for path in (_docx(tmp_path), _xlsx(tmp_path)):
        content_list = asyncio.run(run(path))
        assert hosted[-1] is content_list and len(_images(content_list)) == 1
    assert len(hosted) == 2


def test_office_and_spreadsheet_are_supported_but_legacy_formats_are_not():
    assert {"pdf", "docx", "xlsx"} <= pipeline.SUPPORTED_INGEST_EXTENSIONS
    assert not {"pptx", "doc", "odt", "rtf"} & pipeline.SUPPORTED_INGEST_EXTENSIONS
