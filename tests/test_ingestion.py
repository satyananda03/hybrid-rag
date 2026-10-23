import asyncio
import io
import json
import os
from pathlib import Path
import pymupdf
import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
from hybrid_rag import progress
from hybrid_rag.config import RagSettings, settings
from hybrid_rag.core.chunk_schema import format_location, normalize_page
from hybrid_rag.ingestion.parser import router
from hybrid_rag.ingestion.chunking.parent_child import _page_range, chunking_parent_child
from hybrid_rag.ingestion.chunking.structure import DocumentChunker, SheetChunker, SlideChunker, chunker_for
from hybrid_rag.ingestion.knowledge_graph.entity_repairment import (
    OntologyOrphanRepairOutput,
    OrphanRepairOutput,
    accept_repair,
)
from hybrid_rag.ingestion.knowledge_graph.schema import EntityExtractionOutput, OntologyExtractionOutput
from hybrid_rag.ingestion.doc_status_fields import doc_status_reset_metadata, doc_status_transition_metadata
from hybrid_rag.ingestion.knowledge_graph.chain import get_kg_chains
from hybrid_rag.ingestion.knowledge_graph.ontology_schema import (
    load_ontology_schema,
    prune_chunk,
    render_schema_for_prompt,
    resolve_ontology_path,
)
from hybrid_rag.ingestion.multimodal.content import page_offsets_of, separate_content
from hybrid_rag.ingestion.parser import figures
from hybrid_rag.ingestion.parser.pdf import slides
from hybrid_rag.ingestion.knowledge_graph.entity_resolution import acronym_glossary
from hybrid_rag.ingestion.parser.spreadsheet.parser import parse_xlsx
from hybrid_rag.ingestion.parser.figures import FIGURE_MAX_PIXELS, cap_image_size, save_for_vlm
from hybrid_rag.storage.s3 import objects as s3_objects
from hybrid_rag.storage.s3 import require_bucket, upload_figures_to_s3

def images(items):
    return [i for i in items if i["type"] == "image"]

@pytest.fixture
def xlsx_path(tmp_path, make_png):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    ws["A1"], ws["B1"] = "Nama", "Nilai"
    ws["A2"], ws["B2"] = "x", 1
    ws.add_image(XlImage(str(make_png("small.png", (100, 100)))), "D2")
    ws.add_image(XlImage(str(make_png("big.png", (900, 900)))), "D20")
    path = tmp_path / "a.xlsx"
    wb.save(path)
    return path

def test_xlsx_keeps_only_large_images_and_uses_the_sheet_as_page(xlsx_path, tmp_path):
    [image] = images(parse_xlsx(xlsx_path, tmp_path / "out"))
    assert Image.open(image["img_path"]).size == (900, 900)
    assert image["page_idx"] == "Rekap"

def test_figure_thresholds_come_from_settings(xlsx_path, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "figure_min_width_px", 0)
    monkeypatch.setattr(settings, "figure_min_height_px", 0)
    assert len(images(parse_xlsx(xlsx_path, tmp_path / "out"))) == 2

def test_cap_image_size_downscales_only_images_exceeding_max_pixels():
    small = Image.new("RGB", (800, 600))
    assert cap_image_size(small).size == (800, 600)

    big = Image.new("RGB", (FIGURE_MAX_PIXELS * 2, FIGURE_MAX_PIXELS))
    capped = cap_image_size(big)
    assert max(capped.size) == FIGURE_MAX_PIXELS
    assert capped.size == (FIGURE_MAX_PIXELS, FIGURE_MAX_PIXELS // 2)

def parse(path):
    async def go():
        async with router.parse_file(path) as content_list:
            return content_list
    return asyncio.run(go())

@pytest.fixture
def hosted(monkeypatch):
    calls = []
    monkeypatch.setattr(router, "host_figures", calls.append)
    return calls

def test_supported_extensions_follow_the_routing_table():
    assert router.SUPPORTED_INGEST_EXTENSIONS == frozenset(router.PARSERS)
    assert {"pdf", "docx", "pptx", "doc", "odt", "rtf", "xlsx", "md", "txt", "json"} <= router.SUPPORTED_INGEST_EXTENSIONS

def test_xlsx_is_parsed_natively_then_hosted(xlsx_path, hosted):
    content_list = parse(xlsx_path)
    assert hosted == [content_list] and len(images(content_list)) == 1

@pytest.mark.parametrize(
    "name, page, route",
    [
        ("a.docx", (595, 842), "textract"),  # A4 portrait document
        ("a.pptx", (1440, 810), "slides"),  # 16:9 deck
        ("a.pptx", (1024, 768), "slides"),  # 4:3 deck: is_slide_deck says no, pptx still is a deck
    ],
)
def test_office_files_are_converted_then_take_the_pdf_route(tmp_path, monkeypatch, hosted, name, page, route):
    source = tmp_path / name
    source.write_bytes(b"office")
    outdir = tmp_path / "lo-out"
    seen = {}
    async def fake_convert(path):
        outdir.mkdir()
        return str(make_pdf(outdir / f"{Path(path).stem}.pdf", [page]))
    async def fake_textract(path, output_file):
        seen["textract"] = Path(path)
        return {}
    async def fake_slides(path, tmp_dir, entity_registry=None):
        seen["slides"] = path
        return [{"type": "text", "text": "slide", "page_idx": 0}]
    monkeypatch.setattr(router, "convert_office_to_pdf", fake_convert)
    monkeypatch.setattr(router, "parse_document_to_json", fake_textract)
    monkeypatch.setattr(router, "reconstruct_content_list", lambda raw, output_file: [{"type": "text", "text": "isi"}])
    monkeypatch.setattr(router, "parse_slides", fake_slides)
    content_list = parse(source)
    assert list(seen) == [route] and seen[route].suffix == ".pdf"
    assert seen[route].parent != outdir and not outdir.exists()  # moved into tmp_dir, converter dir removed
    assert hosted == [content_list]

def test_pdf_goes_through_textract_then_hosting_and_cleans_up(tmp_path, monkeypatch, hosted):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    seen = {}
    async def fake_textract(path, output_file):
        seen["path"], seen["tmp_dir"] = path, Path(output_file).parent
        return {"raw": True}
    monkeypatch.setattr(router, "parse_document_to_json", fake_textract)
    monkeypatch.setattr(router, "reconstruct_content_list", lambda raw, output_file: [{"type": "text", "text": "isi"}])
    async def go():
        async with router.parse_file(pdf) as content_list:
            assert seen["tmp_dir"].exists()  # still there while the caller inserts
            return content_list
    assert asyncio.run(go()) == [{"type": "text", "text": "isi", "layout": "document"}]
    assert seen["path"] == str(pdf) and len(hosted) == 1
    assert not seen["tmp_dir"].exists()

def test_text_files_are_not_hosted(tmp_path, hosted):
    md = tmp_path / "a.md"
    md.write_text("halo", encoding="utf-8")
    assert parse(md) == [{"type": "text", "text": "halo", "layout": "document"}] and hosted == []


def test_content_list_images_are_fitted_then_hosted(tmp_path, hosted, monkeypatch):
    fitted = []
    monkeypatch.setattr(router, "fit_figures", lambda content_list, tmp_dir: fitted.append(content_list))
    prebuilt = tmp_path / "a.json"
    prebuilt.write_text(json.dumps([{"type": "image", "img_path": "https://s3/x.jpg"}]), encoding="utf-8")
    content_list = parse(prebuilt)
    assert content_list[0]["img_path"] == "https://s3/x.jpg" and fitted == hosted == [content_list]


def test_parse_file_rejects_bad_input(tmp_path):
    with pytest.raises(FileNotFoundError):
        parse(tmp_path / "missing.pdf")

    unsupported = tmp_path / "a.ppt"
    unsupported.write_bytes(b"x")
    with pytest.raises(ValueError, match=r"\.ppt tidak didukung"):
        parse(unsupported)

    bad = tmp_path / "bad.json"
    bad.write_text('{"not": "a list"}', encoding="utf-8")
    with pytest.raises(ValueError, match="array of object"):
        parse(bad)

class FakeS3Client:
    def __init__(self):
        self.puts = []

    def upload_fileobj(self, fileobj, bucket, key, ExtraArgs):
        if key.endswith("bad.png"):
            raise OSError("boom")
        self.puts.append((bucket, key, ExtraArgs["ContentType"]))

def test_figures_get_urls_or_the_whole_batch_fails(tmp_path, make_png, monkeypatch):
    client = FakeS3Client()
    monkeypatch.setattr(s3_objects, "get_client", lambda: client)
    monkeypatch.setattr(s3_objects, "figure_key", lambda path: f"figs/{Path(path).name}")
    monkeypatch.setattr(settings, "s3_bucket", "bkt")
    monkeypatch.setattr(settings, "s3_region", "ap-southeast-3")
    ok, bad = make_png("ok.png", (10, 10)), make_png("bad.png", (10, 10))

    items = [{"img_path": str(ok)}]
    upload_figures_to_s3(items)
    assert items == [{"img_path": "https://bkt.s3.ap-southeast-3.amazonaws.com/figs/ok.png"}]

    # Satu gagal → error, dan tidak ada img_path yang menunjuk ke file lokal yang akan dihapus.
    items = [{"img_path": str(ok)}, {"img_path": str(bad)}]
    with pytest.raises(RuntimeError, match="bad.png"):
        upload_figures_to_s3(items)
    assert items == [{"img_path": str(ok)}, {"img_path": str(bad)}]

def test_require_bucket_names_the_variable(monkeypatch):
    monkeypatch.setattr(settings, "s3_bucket", None)
    with pytest.raises(ValueError, match="S3_BUCKET"):
        require_bucket()

def node(name, type_, **props):
    return {name: [{"entity_name": name, "entity_type": type_, "properties": props}]}

def edge(src, tgt, type_):
    return {(src, tgt): [{"src_id": src, "tgt_id": tgt, "relation_type": type_, "properties": {}}]}

def test_ontology_loads_and_renders_for_a_prompt_template():
    schema = load_ontology_schema("hukum")
    assert {n.label for n in schema.node_types} >= {"Peraturan", "Pasal", "Lembaga"}
    assert render_schema_for_prompt(schema).format() == str(schema.model_dump(exclude_none=True))

@pytest.mark.parametrize("name", ["../hukum", "nope"])
def test_unknown_or_unsafe_ontology_names_fail(name):
    with pytest.raises((ValueError, FileNotFoundError)):
        resolve_ontology_path(name)

def test_pruning_keeps_the_schema_and_repairs_direction():
    nodes = {
        **node("UU 1/2020", "Peraturan", tahun="2020"),
        **node("Pasal 1", "Pasal"),
        **node("Kucing", "Hewan"), 
    }
    edges = {
        **edge("Pasal 1", "UU 1/2020", "memilikiPasal"),  # reversed: repaired
        **edge("UU 1/2020", "Kucing", "mengaturTentang"),  # endpoint pruned
    }
    kept_nodes, kept_edges = asyncio.run(prune_chunk(nodes, edges, load_ontology_schema("hukum"), "c1"))
    assert set(kept_nodes) == {"UU 1/2020", "Pasal 1"}
    assert kept_nodes["UU 1/2020"][0]["properties"] == {"tahun": "2020"}
    assert list(kept_edges) == [("UU 1/2020", "Pasal 1")]

def test_chains_are_built_per_ontology(fake_llm):
    default, onto = get_kg_chains(), get_kg_chains("hukum")
    assert default.schema is None and get_kg_chains("") is default
    assert onto.schema is not None
    system = onto.extraction_prompt.messages[0].prompt.template
    assert "Peraturan" in system and "{ontology_schema}" not in system
    repair = {"input_text": "x", "heading_context_block": "", "entity_names": "A", "orphan_names": "B"}
    # the repair pass answers in its own contract (FakeLLM builds whatever model the chain is bound to)
    assert type(asyncio.run(default.repair_chain.ainvoke(repair))) is OrphanRepairOutput
    assert type(asyncio.run(onto.repair_chain.ainvoke(repair))) is OntologyOrphanRepairOutput

def test_location_helpers():
    assert [normalize_page(v) for v in (0, 3, " Rekap ", "", None, True, 1.5)] == [0, 3, "Rekap", None, None, None, None]
    assert format_location(None) is None
    assert format_location(2) == format_location(2, 2) == "hal. 3"
    assert format_location(2, 3) == "hal. 3–4"
    assert format_location("Rekap") == 'sheet "Rekap"'
    assert format_location("Rekap", "Detail") == 'sheet "Rekap" – "Detail"'

def test_sheet_pages_flow_into_chunk_ranges():
    offsets = page_offsets_of(["isi rekap", "isi detail"], ["Rekap", "Detail"])
    starts = [o["start"] for o in offsets]
    assert _page_range(offsets, starts, (0, len("isi rekap\n\nisi"))) == {"page_start": "Rekap", "page_end": "Detail"}
    assert page_offsets_of(["teks .md"], [None]) == []

def test_excel_sheets_are_hard_chunk_boundaries():
    """Label sheet tanpa URL tidak boleh bersebelahan dengan URL sheet lain (kasus Banten vs JABAR)."""
    from hybrid_rag.core.tokenizer import get_tokenizer

    content_list = [
        {"type": "text", "text": "Link Data Provinsi BANTEN", "page_idx": "13. BANTEN"},
        {"type": "text", "text": "Link Data Provinsi JAWA BARAT", "page_idx": "14. JABAR"},
        {"type": "table", "table_body": "| No |", "page_idx": "14. JABAR"},
        {"type": "text", "text": "https://drive.google.com/jabar", "page_idx": "14. JABAR"},
    ]
    content, items, options = separate_content(content_list)
    parents, children = chunking_parent_child(get_tokenizer(), content, chunk_options=options)
    for chunk in parents + children:
        assert not ("BANTEN" in chunk["content"] and "jabar" in chunk["content"])
        assert chunk["sidecar"]["page_start"] == chunk["sidecar"]["page_end"]
    assert "BANTEN" not in items[0]["_leading_text"]

    # halaman PDF tetap mengalir: kalimat menyambung antar halaman
    pdf = page_offsets_of(["akhir halaman satu", "awal halaman dua"], [0, 1])
    parents, _ = chunking_parent_child(get_tokenizer(), "akhir halaman satu\n\nawal halaman dua", chunk_options={"page_offsets": pdf})
    assert len(parents) == 1 and parents[0]["sidecar"] == {"page_start": 0, "page_end": 1}

def test_parents_follow_heading_sections():
    """merge_peers: satu section per parent, dipotong di antara paragraf hanya bila lebih besar dari batas."""
    from hybrid_rag.core.tokenizer import get_tokenizer

    long_section = [f"Kalimat ke-{i} menjelaskan arah kebijakan ekonomi kreatif secara rinci." for i in range(8)]
    content_list = [
        {"type": "text", "text": "Judul Dokumen", "text_level": 1, "page_idx": 0},
        {"type": "text", "text": "Pengantar singkat dokumen.", "page_idx": 0},
        {"type": "text", "text": "BAB II", "text_level": 2, "page_idx": 0},
        {"type": "text", "text": "A. Arah", "text_level": 2, "page_idx": 0},
        *({"type": "text", "text": p, "page_idx": 1} for p in long_section),
        {"type": "text", "text": "B. Sasaran", "text_level": 2, "page_idx": 1},
        {"type": "text", "text": "Sasaran singkat.", "page_idx": 1},
    ]
    content, _, options = separate_content(router._with_layout(content_list, "document"))
    tokenizer = get_tokenizer()
    parents, children = chunking_parent_child(tokenizer, content, chunk_options=options, parent_max_tokens=60)

    headings = [(p["heading"]["parent_headings"], p["heading"]["heading"]) for p in parents]
    section_a = [p for p, h in zip(parents, headings) if h == (["Judul Dokumen"], "A. Arah")]
    assert headings[0] == ([], "Judul Dokumen") and headings[-1] == (["Judul Dokumen"], "B. Sasaran")
    assert len(section_a) > 1 and all(p["tokens"] <= 60 for p in parents)
    assert section_a[0]["content"].startswith("BAB II\n\nA. Arah\n\nKalimat ke-0")  # heading kosong ikut section berikutnya
    assert "".join(p["content"] for p in section_a).count("Kalimat ke-") == 8
    assert all("Sasaran" not in p["content"] for p in parents[:-1])
    assert not any("heading" in c for c in children)

def test_markdown_and_text_level_mark_one_heading():
    from hybrid_rag.core.tokenizer import get_tokenizer

    content, _, options = separate_content(router._with_layout([
        {"type": "text", "text": "## Latar Belakang", "text_level": 2},
        {"type": "text", "text": "Isi latar belakang."},
        {"type": "text", "text": "### Rincian"},
        {"type": "text", "text": "Isi rincian."},
    ], "document"))
    spans = chunker_for(get_tokenizer(), options).split(content)
    assert [[h.text for h in s.headings] for s in spans] == [["Latar Belakang"], ["Latar Belakang", "Rincian"]]

def test_untitled_slide_stays_its_own_parent():
    from hybrid_rag.core.tokenizer import get_tokenizer

    content, _, options = separate_content(router._with_layout([
        {"type": "text", "text": "# Agenda\n\nPembukaan acara.", "page_idx": 0},
        {"type": "text", "text": "Lanjutan tanpa judul.", "page_idx": 1},
    ], "slides"))
    parents, _ = chunking_parent_child(get_tokenizer(), content, chunk_options=options)
    assert [p["content"] for p in parents] == ["# Agenda\n\nPembukaan acara.", "Lanjutan tanpa judul."]
    assert parents[0]["heading"]["heading"] == "Agenda" and "heading" not in parents[1]

def textract_page(page_children, layouts):
    """A one-page raw Textract response: ``layouts`` maps block id → (BlockType, text or child ids)."""
    geometry = {"Geometry": {"BoundingBox": {"Left": 0.1, "Top": 0.1, "Width": 0.5, "Height": 0.05}}}
    blocks = [{"Id": "page", "BlockType": "PAGE", **geometry, "Relationships": [{"Type": "CHILD", "Ids": page_children}]}]
    for block_id, (block_type, body) in layouts.items():
        children = body if isinstance(body, list) else [f"{block_id}-line"]
        blocks.append({"Id": block_id, "BlockType": block_type, **geometry, "Relationships": [{"Type": "CHILD", "Ids": children}]})
        if isinstance(body, str):
            blocks.append({"Id": f"{block_id}-line", "BlockType": "LINE", "Text": body, **geometry})
    return {"images_dir": "", "pages": [{"page_idx": 0, "img_w": 100, "img_h": 100, "Blocks": blocks}]}

def test_textract_lists_are_read_once_and_only_the_first_heading_is_a_title():
    from hybrid_rag.ingestion.parser.pdf.amazon_textract import reconstruct_content_list

    layouts = {
        "s1": ("LAYOUT_SECTION_HEADER", "Program"),
        "l1": ("LAYOUT_LIST", ["i1", "i2"]),
        "i1": ("LAYOUT_TEXT", "a. Alamat;"),
        "i2": ("LAYOUT_TEXT", "b. Email."),
        "t1": ("LAYOUT_TITLE", "Kemitraan"),
        "s2": ("LAYOUT_SECTION_HEADER", "Kontak"),
    }
    items = reconstruct_content_list(textract_page(list(layouts), layouts), output_file=None)
    assert [(i["text"], i.get("text_level")) for i in items] == [
        ("Program", 2), ("a. Alamat;", None), ("b. Email.", None), ("Kemitraan", 2), ("Kontak", 2),
    ]

    # list items not on the PAGE are read through their list; a TITLE that comes first stays the title
    layouts = {"t1": ("LAYOUT_TITLE", "Judul"), **{k: layouts[k] for k in ("l1", "i1", "i2")}}
    items = reconstruct_content_list(textract_page(["t1", "l1"], layouts), output_file=None)
    assert [(i["text"], i.get("text_level")) for i in items] == [("Judul", 1), ("a. Alamat; b. Email.", None)]

def test_router_records_layout_on_text_items_only():
    content_list = [{"type": "text", "text": "a"}, {"type": "table", "table_body": "|a|"}, {"type": "text", "text": "b", "layout": "sheet"}]
    router._with_layout(content_list, "document")
    assert [item.get("layout") for item in content_list] == ["document", None, "sheet"]

def test_chunker_for_documents_without_a_recorded_layout():
    from hybrid_rag.core.tokenizer import get_tokenizer

    tokenizer = get_tokenizer()
    assert type(chunker_for(tokenizer, {"page_offsets": [{"page": "Rekap", "start": 0}]})) is SheetChunker
    assert type(chunker_for(tokenizer, {"page_offsets": [{"page": 0, "start": 0, "image": "https://s3/s0.jpg"}]})) is SlideChunker
    assert type(chunker_for(tokenizer, {"page_offsets": [{"page": 0, "start": 0}]})) is DocumentChunker
    assert type(chunker_for(tokenizer, {"layout": "tidak-dikenal"})) is DocumentChunker

class RecordingSink:
    def __init__(self):
        self.events = []
    def stage(self, name):
        self.events.append(("stage", name, progress.current_doc_id.get(), progress.current_workspace.get()))
    def skip(self, name, reason):
        self.events.append(("skip", name, reason))
    def activity(self, message, level="info"):
        self.events.append(("activity", message, level))

def test_progress_without_a_sink_is_a_noop():
    progress.set_sink(None)
    progress.stage("parse"), progress.skip("x", "y"), progress.activity("z")

def test_progress_sink_receives_events_with_the_library_set_context():
    sink = RecordingSink()
    progress.set_sink(sink)
    try:
        async def run():
            progress.current_doc_id.set("doc-1")
            progress.current_workspace.set("ws")
            progress.stage("chunk")
            progress.skip("multimodal", "tidak ada gambar")
            progress.activity("menunggu", level="warning")
        asyncio.run(run())
    finally:
        progress.set_sink(None)
    assert sink.events == [
        ("stage", "chunk", "doc-1", "ws"),
        ("skip", "multimodal", "tidak ada gambar"),
        ("activity", "menunggu", "warning"),
    ]

def test_settings_require_checks_only_the_roles_asked_for():
    s = RagSettings(_env_file=None, KEYWORD_LLM_MODEL="m", KEYWORD_LLM_BINDING_HOST="h")
    s.require()  # no role, nothing to check
    with pytest.raises(RuntimeError) as e:
        s.require("keyword", "contextualize")
    assert str(e.value) == (
        "Missing required env vars: KEYWORD_LLM_BINDING_API_KEY, CONTEXTUALIZE_LLM_MODEL, "
        "CONTEXTUALIZE_LLM_BINDING_HOST, CONTEXTUALIZE_LLM_BINDING_API_KEY"
    )


def test_every_model_role_is_configured_by_the_same_four_settings():
    from hybrid_rag.config import MODEL_ROLES

    for field, env in MODEL_ROLES.values():
        for suffix in ("model", "binding_host", "binding_api_key"):
            assert hasattr(settings, f"{field}_{suffix}"), f"{field}_{suffix}"
        aliases = {f.validation_alias for f in RagSettings.model_fields.values()}
        assert {f"{env}_MODEL", f"{env}_BINDING_HOST", f"{env}_BINDING_API_KEY"} <= aliases

# --- provider contract ---------------------------------------------------------------


@pytest.fixture
def recorded_hooks(monkeypatch):
    """Installs recording hooks for the three model factories; returns their call logs."""
    from hybrid_rag import providers

    calls = {"llm": [], "embeddings": [], "reranker": []}
    monkeypatch.setitem(providers._hooks, "get_llm", lambda *a, **kw: calls["llm"].append((a, kw)))
    monkeypatch.setitem(providers._hooks, "get_embeddings", lambda *a, **kw: calls["embeddings"].append((a, kw)))
    monkeypatch.setitem(providers._hooks, "get_reranker", lambda *a, **kw: calls["reranker"].append((a, kw)))
    monkeypatch.setattr(settings, "max_retries", 7)
    return calls


def test_every_model_factory_gets_the_same_connection_arguments(recorded_hooks):
    from hybrid_rag import providers

    providers.get_llm("m", base_url="h", api_key="k", timeout=9.0, max_tokens=None)
    providers.get_embeddings("m", base_url="h", api_key="k", timeout=9.0)
    providers.get_reranker("m", base_url="h", api_key="k", timeout=9.0)

    connection = {"base_url": "h", "api_key": "k", "timeout": 9.0, "max_retries": 7}
    assert recorded_hooks["llm"] == [(("m",), {**connection, "max_tokens": None})]
    assert recorded_hooks["embeddings"] == [(("m",), connection)]
    assert recorded_hooks["reranker"] == [(("m",), connection)]


def test_vlm_falls_back_to_the_llm_factory_uncapped(recorded_hooks, monkeypatch):
    from hybrid_rag import providers

    monkeypatch.setitem(providers._hooks, "get_vlm", None)
    providers.get_vlm("m", base_url="h", api_key="k", timeout=9.0)
    assert recorded_hooks["llm"] == [(("m",), {"base_url": "h", "api_key": "k", "timeout": 9.0, "max_retries": 7, "max_tokens": None})]


def test_temperature_is_forwarded_only_when_the_library_sets_it(recorded_hooks):
    from hybrid_rag import providers

    providers.get_llm("m", base_url=None, api_key=None)
    providers.get_llm("m", base_url=None, api_key=None, temperature=0.0)
    providers.get_llm("global.xai.grok-4.6", base_url=None, api_key=None, temperature=0.0)  # rejects the field
    assert ["temperature" in kw for _, kw in recorded_hooks["llm"]] == [False, True, False]


def test_reranker_is_none_without_a_host_hook(monkeypatch):
    from hybrid_rag import providers

    monkeypatch.setitem(providers._hooks, "get_reranker", None)
    assert providers.get_reranker("m", base_url="h", api_key="k") is None


def test_embedding_func_builds_its_client_from_settings(recorded_hooks, monkeypatch):
    from hybrid_rag.providers import build_embedding_func

    monkeypatch.setattr(settings, "embedding_model", "emb-model")
    monkeypatch.setattr(settings, "embedding_binding_host", "http://emb")
    monkeypatch.setattr(settings, "embedding_binding_api_key", "emb-key")
    monkeypatch.setattr(settings, "embedding_timeout", 11.0)
    build_embedding_func()
    [(args, kwargs)] = recorded_hooks["embeddings"]
    assert args == ("emb-model",)
    assert kwargs == {"base_url": "http://emb", "api_key": "emb-key", "timeout": 11.0, "max_retries": 7}


def test_env_example_lists_exactly_the_variables_config_reads():
    import re

    read = {f.validation_alias for f in RagSettings.model_fields.values() if isinstance(f.validation_alias, str)}
    listed = re.findall(r"^([A-Z][A-Z0-9_]*)=", Path(".env.example").read_text(), flags=re.M)
    assert sorted(listed) == sorted(read) and len(listed) == len(set(listed))


def test_env_example_changes_no_default():
    example = RagSettings(_env_file=".env.example")
    defaults = RagSettings(_env_file=None)
    # empty values mean "not set": every field must come out as the code default
    assert example.model_dump() == defaults.model_dump()


def test_neighbor_text_is_kept_out_of_modal_extraction():
    from hybrid_rag.ingestion.multimodal.analyze import _NEIGHBOR_TEXT
    from hybrid_rag.ingestion.multimodal.multimodal_template import TEMPLATES

    content = TEMPLATES["image_chunk"].format(
        section_path="S", neighbor_text="halaman\n\npanjang", image_path="https://s3/a.jpg",
        captions="None", footnotes="None", enhanced_caption="Analisis",
    )
    stripped = _NEIGHBOR_TEXT.sub("", content)
    assert "halaman" not in stripped and "Neighbor Text" not in stripped
    assert "- Section Path: S\nImage Path: https://s3/a.jpg\n" in stripped and "Visual Analysis: Analisis" in stripped


# --- slide decks -------------------------------------------------------------------------

def make_pdf(path, sizes):
    doc = pymupdf.open()
    for width, height in sizes:
        doc.new_page(width=width, height=height).insert_text((20, 40), "Timeline Ekraf Kemlu")
    doc.save(path)
    return path


@pytest.mark.parametrize(
    "sizes, deck",
    [
        ([(1440, 810)] * 3, True),  # 16:9
        ([(1280, 800)] * 2, True),  # 16:10
        ([(1440, 810)] * 4 + [(595, 842)], True),  # 80% layar lebar
        ([(1440, 810)] * 3 + [(595, 842)] * 2, False),  # 60%
        ([(595, 842)] * 2, False),  # A4 portrait
        ([(842, 595)] * 2, False),  # A4 landscape
        ([(792, 612)] * 2, False),  # Letter landscape
    ],
)
def test_slide_decks_are_detected_from_page_shape(tmp_path, sizes, deck):
    assert slides.is_slide_deck(make_pdf(tmp_path / "a.pdf", sizes)) is deck


def test_unreadable_pdf_is_left_to_textract(tmp_path):
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    assert slides.is_slide_deck(pdf) is False


class FakeSlideChain:
    def __init__(self, reply=lambda inputs: f"# Slide {inputs['page_number']}\nTimeline Ekraf Kemlu"):
        self.reply, self.calls = reply, []

    async def ainvoke(self, inputs):
        self.calls.append(inputs)
        return self.reply(inputs)


def test_each_slide_becomes_a_text_item_with_its_slide_image(tmp_path, monkeypatch):
    chain = FakeSlideChain()
    monkeypatch.setattr(slides, "slide_transcription_chain", lambda: chain)
    items = asyncio.run(slides.parse_slides(make_pdf(tmp_path / "a.pdf", [(1440, 810)] * 2), tmp_path))
    assert [(i["type"], i["page_idx"], i["text"].splitlines()[0]) for i in items] == [
        ("text", 0, "# Slide 1"), ("text", 1, "# Slide 2"),
    ]
    assert all(Path(i["img_path"]).is_file() for i in items)
    assert chain.calls[0]["reference_text"] == "Timeline Ekraf Kemlu"
    assert chain.calls[0]["style_groups"] == ""  # plain black text: no style block
    assert chain.calls[0]["image_url"].startswith("data:image/jpeg;base64,")


def test_style_groups_group_coloured_and_bold_text_from_the_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=1440, height=810)
    page.insert_text((20, 40), "4 SUBSEKTOR BARU", color=(0.07, 0.39, 0.92))
    page.insert_text((20, 80), "Kuliner", fontname="hebo")
    page.insert_text((20, 120), "Seni Rupa")
    page.insert_text((20, 160), "Sulih Suara", color=(0.07, 0.39, 0.92))
    page.insert_text((20, 200), "Catatan abu-abu", color=(0.5, 0.5, 0.5))
    block = slides.style_groups(page)
    assert block.startswith("<style_groups>\n") and block.endswith("\n</style_groups>")
    assert "4 SUBSEKTOR BARU | Sulih Suara" in block.splitlines()[1]  # same colour → one line
    assert block.splitlines()[1].startswith("warna #")
    assert "tebal: Kuliner" in block
    assert "Seni Rupa" not in block and "abu-abu" not in block


def test_slide_prompt_renders_with_the_inputs_parse_slides_sends():
    from hybrid_rag.ingestion.multimodal.chain import _chat_prompt

    messages = _chat_prompt("slide_transcription", with_image=True).format_messages(
        page_number=3, reference_text="Teks", style_groups="<style_groups>\ntebal: A\n</style_groups>",
        glossary=acronym_glossary("iccn", "ICCN"), image_url="data:image/jpeg;base64,AA==",
    )
    text = messages[1].content[0]["text"]
    assert "Transcribe slide 3" in text and "tebal: A" in text and "ICCN = Indonesia" in text and "{" not in text


def test_slides_get_the_acronyms_their_text_layer_mentions(tmp_path, monkeypatch):
    doc = pymupdf.open()
    doc.new_page(width=1440, height=810).insert_text((20, 40), "DATA JEJARING ICCN")
    doc.new_page(width=1440, height=810).insert_text((20, 40), "Timeline Ekraf")
    doc.save(tmp_path / "a.pdf")
    chain = FakeSlideChain()
    monkeypatch.setattr(slides, "slide_transcription_chain", lambda: chain)
    asyncio.run(slides.parse_slides(tmp_path / "a.pdf", tmp_path, "iccn"))
    by_page = {c["page_number"]: c["glossary"] for c in chain.calls}
    assert "ICCN = Indonesia Creative Cities Network" in by_page[1] and by_page[2] == ""


def test_contextualize_gets_the_acronyms_of_the_parent(monkeypatch):
    from hybrid_rag.ingestion.chunking import parent_child

    seen = []
    class Chain:
        async def ainvoke(self, inputs):
            seen.append(inputs["glossary"])
            return "konteks"
    monkeypatch.setattr(parent_child, "contextualize_chain", lambda: Chain())
    parents = [{"content": "DATA JEJARING ICCN BANTEN"}, {"content": "tanpa akronim"}]
    children = [{"content": "a", "parent_index": 0, "chunk_order_index": 0},
                {"content": "b", "parent_index": 1, "chunk_order_index": 1}]
    asyncio.run(parent_child.contextualize_children(parents, children, entity_registry="iccn"))
    assert "ICCN = Indonesia Creative Cities Network" in seen[0] and seen[1] == ""
    asyncio.run(parent_child.contextualize_children(parents, children))
    assert seen[2:] == ["", ""]


def test_table_and_image_context_carries_the_document_acronyms():
    from hybrid_rag.ingestion.multimodal.base_processor import BaseModalProcessor

    block = BaseModalProcessor._context_block(None, {"table_body": "| ICCN | 3 |", "_entity_registry": "iccn"})
    assert block == acronym_glossary("iccn", "ICCN") and block
    assert BaseModalProcessor._context_block(None, {"table_body": "| ICCN | 3 |"}) == ""


def test_summary_prompt_renders_with_and_without_a_glossary():
    from hybrid_rag.ingestion.knowledge_graph.chain import _prompt

    template = _prompt("summarize_entity_descriptions")["user_prompt"]
    args = dict(description_type="Entity", description_name="ICCN", description_list="{}", summary_length=100)
    assert "ICCN = Indonesia" in template.format(**args, glossary=acronym_glossary("iccn", "ICCN"))
    assert "{" not in template.format(**args, glossary="").replace('{}', '')


def test_parse_file_rejects_an_unknown_entity_registry(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("ICCN", encoding="utf-8")
    async def go():
        async with router.parse_file(md, entity_registry="tidak-ada"):
            pass
    with pytest.raises(FileNotFoundError, match="iccn"):
        asyncio.run(go())


def test_an_empty_transcription_fails_the_document(tmp_path, monkeypatch):
    monkeypatch.setattr(slides, "slide_transcription_chain", lambda: FakeSlideChain(lambda inputs: "  "))
    with pytest.raises(RuntimeError, match="slide 1"):
        asyncio.run(slides.parse_slides(make_pdf(tmp_path / "a.pdf", [(1440, 810)]), tmp_path))


def test_coverage_counts_text_layer_words_in_the_transcript():
    assert slides._coverage("Akhir Juli diajukan PTRI", "# Timeline\n- Akhir Juli: diajukan oleh PTRI NY") == 1.0
    assert slides._coverage("Akhir Juli diajukan PTRI", "- Akhir Juli") == 0.5
    assert slides._coverage("", "apa saja") == 1.0


def test_slide_decks_skip_textract_and_are_hosted(tmp_path, monkeypatch, hosted):
    async def fake_textract(path, output_file):
        raise AssertionError("dek slide tidak boleh ke Textract")
    async def fake_slides(path, tmp_dir, entity_registry=None):
        return [{"type": "text", "text": "slide", "page_idx": 0, "img_path": str(tmp_dir / "slide_0.jpg")}]
    monkeypatch.setattr(router, "parse_document_to_json", fake_textract)
    monkeypatch.setattr(router, "parse_slides", fake_slides)
    content_list = parse(make_pdf(tmp_path / "deck.pdf", [(1440, 810)]))
    assert [i["text"] for i in content_list] == ["slide"] and hosted == [content_list]


def test_host_figures_uploads_exactly_the_items_with_an_image(monkeypatch):
    uploaded = []
    monkeypatch.setattr(figures, "upload_figures_to_s3", uploaded.append)
    content_list = [
        {"type": "text", "text": "paragraf"},
        {"type": "image", "img_path": "/tmp/fig.jpg"},
        {"type": "table", "table_body": "|a|"},
        {"type": "text", "text": "slide", "img_path": "/tmp/slide_0.jpg"},
        {"type": "image", "img_path": "https://s3/sudah_di_host.jpg"},
    ]
    figures.host_figures(content_list)
    assert uploaded == [[content_list[1], content_list[3]]]


def test_slide_images_reach_page_offsets_and_sidecars():
    _, _, options = separate_content([
        {"type": "text", "text": "slide satu", "page_idx": 0, "img_path": "https://s3/s0.jpg"},
        {"type": "text", "text": "slide dua", "page_idx": 1, "img_path": "/tmp/upload_gagal.jpg"},
    ])
    offsets = options["page_offsets"]
    assert offsets == [{"page": 0, "start": 0, "image": "https://s3/s0.jpg"}, {"page": 1, "start": 12}]
    starts = [o["start"] for o in offsets]
    assert _page_range(offsets, starts, (0, 16)) == {"page_start": 0, "page_end": 1, "images": ["https://s3/s0.jpg"]}
    assert _page_range(offsets, starts, (12, 16)) == {"page_start": 1, "page_end": 1}


def test_documents_without_slides_keep_their_offsets_and_sidecars():
    text, items, options = separate_content([
        {"type": "text", "text": "isi halaman", "page_idx": 0},
        {"type": "image", "img_path": "https://s3/fig.jpg", "page_idx": 0},
        {"type": "table", "table_body": "|a|", "page_idx": 1},
        {"type": "text", "text": "isi berikutnya", "page_idx": 1},
    ])
    assert options == {"page_offsets": [{"page": 0, "start": 0}, {"page": 1, "start": 13}]}
    assert len(items) == 2 and "https://s3/fig.jpg" not in text
    assert _page_range(options["page_offsets"], [0, 13], (0, 20)) == {"page_start": 0, "page_end": 1}


# --- images sent to the VLM -------------------------------------------------------------

def png_bytes(img):
    buffer = io.BytesIO()
    img.save(buffer, "PNG")
    return buffer.getvalue()


def test_an_embedded_image_that_fits_is_kept_byte_for_byte(tmp_path):
    original = png_bytes(Image.new("RGB", (900, 600), "red"))
    with Image.open(io.BytesIO(original)) as img:
        path = save_for_vlm(img, tmp_path / "a.png", original=original)
    assert path == tmp_path / "a.png" and path.read_bytes() == original


def test_an_oversized_image_is_scaled_to_the_pixel_cap(tmp_path):
    path = save_for_vlm(Image.new("RGB", (FIGURE_MAX_PIXELS * 2, FIGURE_MAX_PIXELS), "red"), tmp_path / "a.png")
    assert path.suffix == ".jpg"
    with Image.open(path) as saved:
        assert saved.size == (FIGURE_MAX_PIXELS, FIGURE_MAX_PIXELS // 2)


def test_an_image_over_the_byte_budget_is_reencoded_until_it_fits(tmp_path, monkeypatch):
    from hybrid_rag.ingestion.parser import figures as figs
    monkeypatch.setattr(figs, "FIGURE_MAX_BYTES", 60_000)
    noise = Image.frombytes("RGB", (800, 800), os.urandom(800 * 800 * 3))  # tidak bisa dikompres
    original = png_bytes(noise)
    path = save_for_vlm(noise, tmp_path / "a.png", original=original)
    assert path.suffix == ".jpg" and path.stat().st_size <= 60_000
    with Image.open(path) as saved:
        assert max(saved.size) < 800  # kualitas saja tidak cukup, ukurannya ikut turun


def test_transparent_areas_become_white(tmp_path):
    clear = Image.new("RGBA", (FIGURE_MAX_PIXELS + 10, 20), (0, 0, 0, 0))
    with Image.open(save_for_vlm(clear, tmp_path / "a.png")) as saved:
        assert saved.convert("RGB").getpixel((5, 5)) == (255, 255, 255)


class FakeResponse:
    def __init__(self, content):
        self.content = content

    def raise_for_status(self):
        pass


def test_ready_made_images_are_fitted_only_when_over_the_limits(tmp_path, monkeypatch):
    small = png_bytes(Image.new("RGB", (800, 600), "red"))
    big = png_bytes(Image.new("RGB", (FIGURE_MAX_PIXELS * 2, 100), "red"))
    local = tmp_path / "local.png"
    local.write_bytes(small)
    remote = {"https://s3/small.png": small, "https://s3/big.png": big}
    monkeypatch.setattr(figures.httpx, "get", lambda url, **kwargs: FakeResponse(remote[url]))
    content_list = [
        {"type": "text", "text": "paragraf"},
        {"type": "image", "img_path": "https://s3/small.png"},
        {"type": "image", "img_path": "https://s3/big.png"},
        {"type": "image", "img_path": str(local)},
    ]
    figures.fit_figures(content_list, tmp_path)
    assert content_list[1]["img_path"] == "https://s3/small.png"  # URL permanen yang muat tidak diunggah ulang
    assert content_list[2]["img_path"] == str(tmp_path / "figure_2.jpg")  # di-encode ulang untuk host_figures
    with Image.open(content_list[2]["img_path"]) as fitted:
        assert max(fitted.size) == FIGURE_MAX_PIXELS
    assert content_list[3]["img_path"] == str(local)


def test_an_unreadable_ready_made_image_fails_the_document(tmp_path, monkeypatch):
    monkeypatch.setattr(figures.httpx, "get", lambda url, **kwargs: FakeResponse(b"bukan gambar"))
    with pytest.raises(RuntimeError, match="https://s3/rusak.png"):
        figures.fit_figures([{"type": "image", "img_path": "https://s3/rusak.png"}], tmp_path)


# --- tokenizer --------------------------------------------------------------------------

SAMPLE = (
    "Konferensi internasional Creative Economy terbesar di dunia. WCCE ke-5 di Jakarta "
    "(21-23 Oktober 2026), dihadiri 50+ negara. Dokumen Jakarta Vision on Creative Economy "
    "masuk menjadi resolusi PBB sebagai dasar menjadikan Ekonomi Kreatif agenda pasca SDGs."
)


def test_the_tokenizer_never_downloads(monkeypatch):
    import requests
    from hybrid_rag.core.tokenizer import TiktokenTokenizer
    monkeypatch.delenv("TIKTOKEN_CACHE_DIR", raising=False)
    monkeypatch.delenv("DATA_GYM_CACHE_DIR", raising=False)
    monkeypatch.setattr(requests, "get", lambda *a, **k: pytest.fail("tiktoken mencoba mengunduh"))
    assert TiktokenTokenizer().encode(SAMPLE)


def test_the_shared_tokenizer_counts_exactly_like_tiktoken():
    import tiktoken
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from hybrid_rag.core.tokenizer import get_tokenizer
    from hybrid_rag.ingestion.multimodal import content
    assert get_tokenizer() is get_tokenizer()
    assert get_tokenizer().encode(SAMPLE) == tiktoken.encoding_for_model("gpt-4o-mini").encode(SAMPLE)
    reference = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
        model_name="gpt-4o-mini", chunk_size=20, chunk_overlap=0, separators=content._SPLIT_SEPARATORS,
    )
    assert content._splitter(20).split_text(SAMPLE) == reference.split_text(SAMPLE)


def test_normalize_string_list_drops_non_strings_without_crashing():
    from hybrid_rag.core.utils import normalize_string_list

    assert normalize_string_list(["a", 1, "", None, "b"], context="t") == ["a", "b"]
    assert normalize_string_list("not a list") == []


def test_ingestion_rag_keeps_purge_and_delete_methods():
    # The resume path (pipeline.py) and delete_document call these by name; no test runs either path.
    from hybrid_rag.ingestion.rag import IngestionRAG

    for name in ("_chunk_ids_from_storage", "_purge_doc_chunks_and_kg", "adelete_by_doc_id", "_get_existing_llm_cache_ids"):
        assert callable(getattr(IngestionRAG, name)), name


def test_accept_repair_keeps_only_known_orphan_edges():
    nodes = {"A": ["a"], "B": ["b"], "O": ["o"], "P": ["p"]}
    edges = {("A", "B"): ["ab"], ("P", "A"): ["pa"]}
    repair = {
        ("O", "A"): ["oa"],    # orphan -> known: kept
        ("O", "X"): ["ox"],    # X is not a node of the chunk: would become an UNKNOWN node
        ("B", "A"): ["ba"],    # known <-> known, neither orphan
        ("A", "P"): ["ap"],    # reverse of an existing edge
    }
    assert accept_repair(nodes, edges, {"O", "P"}, repair) == {("O", "A"): ["oa"]}


@pytest.mark.parametrize("driver", ["sync", "async"])
def test_embed_by_index_pairs_vectors_by_index_not_position(driver):
    from hybrid_rag import providers

    calls = []

    def respond(input):
        calls.append(list(input))
        if len(input) == 1:
            return {"data": [{"index": 0, "embedding": [f"v:{input[0]}"]}]}
        # what the caching proxy returned: out of order, one input twice, one missing, one out of range
        return {"data": [{"index": i, "embedding": [f"v:{input[i] if i < len(input) else '?'}"]} for i in (1, 0, 1, 3, 9)]}

    async def acreate(input):
        return respond(input)

    texts = ["a", "b", "c", "d", "e"]
    if driver == "sync":
        vectors = providers.embed_by_index_sync(respond, texts, batch_size=4)
    else:
        vectors = asyncio.run(providers.embed_by_index(acreate, texts, batch_size=4))
    assert vectors == [[f"v:{t}"] for t in "abcde"]
    assert calls == [["a", "b", "c", "d"], ["e"], ["c"]]  # only the input left without a vector is re-asked
    assert providers.embed_by_index_sync(respond, [], batch_size=4) == []


def test_repair_contract_hides_entities_from_the_model_but_keeps_the_extraction_shape():
    for repair in (OrphanRepairOutput, OntologyOrphanRepairOutput):
        schema = repair.model_json_schema()
        assert list(schema["properties"]) == ["relationships"] and schema["required"] == ["relationships"]
    assert "entities" in EntityExtractionOutput.model_json_schema()["properties"]
    assert issubclass(OntologyOrphanRepairOutput, OntologyExtractionOutput)  # parsed as ontology output

    good = {"source": "A", "target": "B", "keywords": ["k1", "k2"], "description": "A berelasi dengan B."}
    reply = OrphanRepairOutput.model_validate_json(json.dumps({"relationships": [good, {"source": "A"}]}))
    assert reply.entities == []  # absent from the reply, present on the class for the parser and the cache
    assert [(r.source, r.keywords) for r in reply.relationships] == [("A", "k1, k2")]  # the extraction validator still runs
    assert json.loads(reply.model_dump_json())["entities"] == []


def test_repair_request_sends_relationships_only(monkeypatch):
    import httpx
    from langchain_openai import ChatOpenAI

    from hybrid_rag import providers
    from hybrid_rag.ingestion.knowledge_graph import chain

    requests = []

    def handler(request):
        requests.append(json.loads(request.content))
        asked = requests[-1]["response_format"]["json_schema"]["schema"]["properties"]
        message = {"role": "assistant", "content": json.dumps({name: [] for name in asked})}  # an empty answer in the asked shape
        return httpx.Response(200, json={"id": "x", "object": "chat.completion", "created": 0, "model": "m",
                                         "choices": [{"index": 0, "message": message, "finish_reason": "stop"}]})

    llm = ChatOpenAI(model="m", api_key="k", base_url="http://llm.test/v1", max_retries=0,
                     http_async_client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    monkeypatch.setitem(providers._hooks, "get_llm", lambda *_, **__: llm)
    chain._build_chains.cache_clear()
    chain._extract_llm.cache_clear()
    try:
        inputs = {"input_text": "teks", "heading_context_block": "", "entity_names": "A", "orphan_names": "B"}
        reply = asyncio.run(get_kg_chains().repair_chain.ainvoke(inputs))
        asyncio.run(get_kg_chains().extraction_chain.ainvoke({"input_text": "teks", "heading_context_block": ""}))
    finally:
        chain._build_chains.cache_clear()
        chain._extract_llm.cache_clear()
    repair_request, extraction_request = requests
    sent = lambda request: request["response_format"]["json_schema"]["schema"]["properties"]
    assert list(sent(repair_request)) == ["relationships"]  # what the model is asked for
    assert list(sent(extraction_request)) == ["entities", "relationships"]  # the extraction contract is untouched
    assert type(reply) is OrphanRepairOutput and reply.entities == []


def test_ingest_flags_survive_status_transitions_and_resets():
    doc = {"metadata": {"orphan_repair": False, "prune_isolated": False, "chunk_opts": "stale"}}
    for payload in (doc_status_transition_metadata(doc), doc_status_reset_metadata(doc)):
        assert payload["orphan_repair"] is False and payload["prune_isolated"] is False  # False is kept, not "unset"


@pytest.mark.parametrize("orphan_repair", [True, False])
def test_orphan_repair_follows_the_ingest_flag(fake_llm, monkeypatch, orphan_repair):
    from hybrid_rag.ingestion.knowledge_graph import extraction
    from hybrid_rag.ingestion.knowledge_graph.schema import ExtractedEntity

    passes = []
    unlinked = EntityExtractionOutput(entities=[
        ExtractedEntity(name="Kota A", type="Organization", description="Kota A disebut."),
        ExtractedEntity(name="Kota B", type="Organization", description="Kota B disebut."),
    ], relationships=[])

    async def extract(chain, inputs, config, chunk_key, pass_name):
        passes.append(pass_name)
        return unlinked if pass_name == "extraction" else None

    monkeypatch.setattr(extraction, "_extract_or_none", extract)
    chunk = {"c1": {"content": "Kota A dan Kota B.", "file_path": "f", "full_doc_id": "d", "tokens": 5, "chunk_order_index": 0}}
    asyncio.run(extraction.extract_entities(chunk, {"entity_extract_max_gleaning": 0, "orphan_repair": orphan_repair}))
    assert ("orphan repair" in passes) is orphan_repair
